"""The SQL result with no row cap, and its actions pinned above it.

Two things were asked for together. The result was capped at 5,000 rows,
and Copy / CSV / Excel exported what the pane held, so every export of a
bigger answer was silently cut at 5,000 too. And the bar carrying those
actions sat inside the result's scroller, so it scrolled away the moment
the analyst went looking at the rows they wanted to export.

Lifting the cap means a result can be millions of rows. The server holds
them (winnow/sqlresults.py) and the pane pages them in, so most rows of a
big result were never in the browser at all — which is why the exports,
the sort and the tags are asserted here on rows past the first page. The
table is virtualized like the grid, and a few of these tests pin the ways
that goes wrong: rows past the browser's element-height ceiling
unreachable (the grid's own MAX_SPACER_PX story), and columns that change
width as you scroll because auto layout only ever sees the rows that
happen to be painted.

Rows are generated with a recursive CTE, so the shared case needs no
bigger table than the one every other module uses."""

from __future__ import annotations

import io
from pathlib import Path

import pytest

pytestmark = pytest.mark.ui


def _count(n, extra=""):
    return (f"WITH RECURSIVE n(i) AS (SELECT 1 UNION ALL SELECT i + 1 FROM n WHERE i < {n}) "
            f"SELECT i{extra} FROM n")


def _run(page, sql):
    page.click("#tabSql")
    page.wait_for_selector("#sqlview:not([hidden])")
    page.wait_for_function(
        "() => __winnow.S.sqlTabs.length > 0 && !document.getElementById('sqlText').disabled")
    page.locator("#sqlText").fill(sql)
    page.click("#btnRunSql")
    page.wait_for_selector("#sqlResult table", timeout=60_000)


def _scroll_result(page, wheel_settled, dy=0, dx=0):
    """Wheel over the result's rows, the way an analyst scrolls it — not by
    setting scrollTop on whichever element the implementation happens to
    scroll — and let the virtualized window repaint. The default is "all
    the way down"."""
    if not dx and not dy:
        dy = 10 ** 8
    # Measured in the page, in one call, on the scroller rather than a row
    # cell: the rows are rebuilt whenever the pane is laid out or a page of
    # them lands, so a cell looked up in one round trip and measured in the
    # next could already be gone, and bounding_box() answered None.
    at = page.evaluate("() => { const b = document.querySelector('#sqlResult .sql-result-scroll')"
                       ".getBoundingClientRect(); return [b.x + b.width / 2, b.y + b.height / 2]; }")
    page.mouse.move(*at)
    with wheel_settled(page, "#sqlResult .sql-result-scroll"):
        page.mouse.wheel(dx, dy)
    # The rows the scroll brought on screen come from the server.
    page.wait_for_function("() => !document.querySelector('#sqlResult tr.sql-row-pending')", timeout=30_000)


def _hit(page, locator):
    """True when the element is what's actually under its own centre —
    on screen and not scrolled out of its pane or painted over."""
    box = locator.bounding_box()
    return page.evaluate("""([x, y, label]) => {
        const at = document.elementFromPoint(x, y);
        return !!at && !!at.closest('.btn') && at.closest('.btn').textContent.trim() === label;
    }""", [box["x"] + box["width"] / 2, box["y"] + box["height"] / 2, locator.inner_text().strip()])


def test_a_result_past_5000_rows_comes_back_whole(page):
    _run(page, _count(12_000))
    status = page.locator(".sql-result-bar .note-status").inner_text()
    assert "12,000 rows" in status
    assert "truncated" not in status


def test_the_actions_stay_in_view_while_the_result_scrolls(page, wheel_settled):
    """The symptom: scroll down a result and Copy/CSV/Excel/Save are gone."""
    _run(page, _count(12_000, ", printf('%080d', i) AS wide"))
    for label in ("Copy", "CSV…", "Excel…", "Save as table…"):
        assert _hit(page, page.locator(".sql-result-actions .btn", has_text=label)), label
    _scroll_result(page, wheel_settled)
    _scroll_result(page, wheel_settled, dx=10_000)
    for label in ("Copy", "CSV…", "Excel…", "Save as table…"):
        assert _hit(page, page.locator(".sql-result-actions .btn", has_text=label)), \
            f"{label} scrolled out of view with the rows"
    # And the header row is pinned too — it always was; the move to an
    # inner scroller must not have cost it.
    head = page.locator("#sqlResult th").first.bounding_box()
    bar = page.locator(".sql-result-bar").bounding_box()
    assert 0 <= head["y"] - (bar["y"] + bar["height"]) < 12, "the header row sits right under the bar"


def test_only_the_visible_window_is_in_the_dom(page, wheel_settled):
    _run(page, _count(12_000))
    assert page.locator("#sqlResult tr").count() < 300
    _scroll_result(page, wheel_settled)
    last = page.locator("#sqlResult tr").last.locator("td").first.inner_text()
    assert last == "12000", "the last row must be reachable by scrolling"


def test_the_last_of_two_million_rows_is_reachable(page, wheel_settled):
    """Two million rows is ~40M px of table at 20px a row, past the height
    Blink will lay an element out at (33,554,365px). Without the grid's
    spacer cap and rescale the tail of the result is simply not there to
    scroll to — a 2.46M-row table once showed only its first 1.4M."""
    _run(page, _count(2_000_000))
    assert "2,000,000 rows" in page.locator(".sql-result-bar .note-status").inner_text()
    _scroll_result(page, wheel_settled)
    cells = [c for c in page.locator("#sqlResult td:first-child").all_inner_texts() if c]
    assert cells[-1] == "2000000"
    # The last row is ON SCREEN, not merely in the DOM below it.
    row = page.locator("#sqlResult tr").last.bounding_box()
    pane = page.locator("#sqlResult").bounding_box()
    assert row["y"] + row["height"] <= pane["y"] + pane["height"] + 1


def test_columns_keep_their_width_while_scrolling(page, wheel_settled):
    """Auto table layout sizes columns to the rows in the DOM, so a long
    value near the bottom would widen its column only once scrolled to.
    Widths come from every row up front instead."""
    _run(page, _count(12_000, ", CASE WHEN i > 11500 THEN printf('%060d', i) ELSE 'x' END AS v"))
    widths = lambda: page.evaluate(  # noqa: E731
        "() => [...document.querySelectorAll('#sqlResult th')].map((th) => th.getBoundingClientRect().width)")
    before = widths()
    _scroll_result(page, wheel_settled)
    after = widths()
    assert after == pytest.approx(before, abs=1)


def test_copy_takes_every_row(page):
    _run(page, _count(7_000))
    # The clipboard outlives the browser context, so whatever an earlier
    # test copied is still there: wait for it to CHANGE, not to be non-empty.
    page.evaluate("() => navigator.clipboard.writeText('-')")
    page.locator(".sql-result-actions .btn", has_text="Copy").click()
    text = "-"
    for _ in range(50):   # polled with evaluate, which awaits — see test_ui_test_hygiene
        text = page.evaluate("() => navigator.clipboard.readText()")
        if text != "-":
            break
        page.wait_for_timeout(100)
    lines = text.split("\n")
    assert lines[0] == "i"
    assert len(lines) == 7_001, "header + every row, not the first 5,000"
    assert lines[-1] == "7000"


def test_excel_takes_every_row(page):
    from openpyxl import load_workbook

    _run(page, _count(7_000))
    page.locator(".sql-result-actions .btn", has_text="Excel").click()
    page.wait_for_selector(".confirm-overlay input")
    page.locator(".confirm-overlay input").fill("all rows")
    with page.expect_download() as dl:
        page.locator(".confirm-card .btn", has_text="Save").first.click()
    ws = load_workbook(io.BytesIO(Path(dl.value.path()).read_bytes())).active
    assert ws.max_row == 7_001
    assert ws.cell(row=7_001, column=1).value == 7000


def test_csv_takes_every_row(page):
    _run(page, _count(7_000))
    page.locator(".sql-result-actions .btn", has_text="CSV").click()
    page.wait_for_selector(".confirm-overlay input")
    page.locator(".confirm-overlay input").fill("all rows")
    with page.expect_download() as dl:
        page.locator(".confirm-card .btn", has_text="Save").first.click()
    lines = Path(dl.value.path()).read_text(encoding="utf-8-sig").strip().splitlines()
    assert len(lines) == 7_001


def test_a_click_sort_orders_the_whole_result(page):
    _run(page, _count(7_000))
    top = page.locator("#sqlResult td").first.inner_text()
    page.locator("#sqlResult th", has_text="i").click()
    page.locator("#sqlResult th", has_text="i").click()   # descending
    # The sort runs on the server; wait for the rows it sends back.
    page.wait_for_function("(top) => !document.querySelector('#sqlResult tr.sql-row-pending')"
                           " && document.querySelector('#sqlResult td').textContent !== top", arg=top)
    first = page.locator("#sqlResult td").first.inner_text()
    assert first == "7000", "the sort must cover rows past the old cap"


def test_a_tagged_row_past_the_first_page_shows_its_tag(page, wheel_settled):
    """The Tags column used to be joined in the browser over the rows it
    held, so with rows held on the server it must come from the server —
    and a row the pane only fetched while scrolling has to arrive with its
    chip. 200 rids × 40 copies, ordered by rid, puts rid 200 in the last
    40 rows: positions 7,960–7,999, far past the first page."""
    tag = page.evaluate("() => __winnow.S.tags[0]")
    page.evaluate("(t) => __winnow.post('/api/row_tags', { pairs: [[1, 200]], tag_id: t, on: true })", tag["id"])
    try:
        _run(page, "SELECT x.rid FROM src_1 x, (WITH RECURSIVE n(i) AS (SELECT 1 UNION ALL "
                   "SELECT i + 1 FROM n WHERE i < 40) SELECT i FROM n) ORDER BY x.rid, i")
        assert "8,000 rows" in page.locator(".sql-result-bar .note-status").inner_text()
        _scroll_result(page, wheel_settled)
        page.wait_for_function("() => document.querySelector('#sqlResult tbody tr:last-child td')"
                               "?.textContent === '200'")
        last = page.locator("#sqlResult tbody tr[data-i]").last
        assert last.locator(".sql-tag-chip").all_inner_texts() == [tag["name"]]
        # And a row with no tag has no chip: the join is per row, not "the
        # table has tags".
        page.evaluate("() => { document.querySelector('.sql-result-scroll').scrollTop = 0; }")
        page.wait_for_function("() => document.querySelector('#sqlResult tbody tr[data-i=\"0\"]')")
        assert page.locator('#sqlResult tbody tr[data-i="0"] .sql-tag-chip').count() == 0
    finally:
        page.evaluate("(t) => __winnow.post('/api/row_tags', { pairs: [[1, 200]], tag_id: t, on: false })", tag["id"])
