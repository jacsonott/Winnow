"""The SQL result bar's actions: Copy, CSV…, Excel…, Save as table.

The bar was rebuilt because the buttons used to be appended straight
after the status text, with no grouping, so they collided with it and
read as part of the sentence.

Copy, CSV and Excel all export WHAT IS ON SCREEN — the preview, in the
order it is sorted in. Save as table is the one that re-runs the query
in full. That split is deliberate and is what the Excel cases below
pin: a workbook that quietly disagreed with the row count printed above
it would be the worst of the four to get wrong, because it is the one
that leaves the building."""

from __future__ import annotations

import pytest

pytestmark = pytest.mark.ui


def _run(page, sql):
    page.click("#tabSql")
    page.wait_for_selector("#sqlview:not([hidden])")
    page.wait_for_function(
        "() => __winnow.S.sqlTabs.length > 0 && !document.getElementById('sqlText').disabled")
    page.locator("#sqlText").fill(sql)
    page.click("#btnRunSql")
    page.wait_for_selector("#sqlResult table")


def test_actions_are_grouped_to_the_right_of_the_status(page):
    _run(page, "SELECT rid, EventId FROM src_1 ORDER BY rid LIMIT 5")
    bar = page.locator(".sql-result-bar")
    status = bar.locator(".note-status").bounding_box()
    acts = bar.locator(".sql-result-actions").bounding_box()
    # The regression this replaces: buttons butted up against the text.
    assert acts["x"] > status["x"] + status["width"], "actions must not overlap the status text"
    assert [b.strip() for b in page.locator(".sql-result-actions .btn").all_inner_texts()] \
        == ["Copy", "CSV…", "Excel…", "Save as table…"]


def test_copy_takes_the_whole_result_with_headers(page):
    _run(page, "SELECT rid, EventId FROM src_1 ORDER BY rid LIMIT 5")
    page.locator(".sql-result-actions .btn", has_text="Copy").click()
    page.wait_for_timeout(200)
    clip = page.evaluate("() => navigator.clipboard.readText()")
    lines = clip.split("\n")
    assert lines[0] == "rid\tEventId", "a header row is the point of the feature"
    assert len(lines) == 6  # header + 5 rows


def test_copy_prefers_the_selection_when_there_is_one(page):
    _run(page, "SELECT rid, EventId FROM src_1 ORDER BY rid LIMIT 10")
    rows = page.locator("#sqlResult tr")
    rows.nth(1).click()
    rows.nth(3).click(modifiers=["Control"])
    page.wait_for_timeout(150)
    page.locator(".sql-result-actions .btn", has_text="Copy").click()
    page.wait_for_timeout(200)
    lines = page.evaluate("() => navigator.clipboard.readText()").split("\n")
    assert lines[0] == "rid\tEventId"
    assert len(lines) == 3, "header + the two selected rows only"
    page.keyboard.press("Escape")


def test_csv_prompts_with_the_tab_name_as_the_default(page):
    _run(page, "SELECT rid FROM src_1 LIMIT 3")
    page.locator(".sql-result-actions .btn", has_text="CSV").click()
    page.wait_for_selector(".confirm-overlay input")
    # Read the tab's name rather than hard-coding "Query 1": SQL tabs live
    # in the case file and the UI suite shares one, so a rename in another
    # module carries over. The contract is "whatever this tab is called",
    # not one particular string — and definitely not "query-results.csv",
    # which made every tab's export collide in the downloads folder.
    tab_name = page.evaluate("() => __winnow.activeSqlTab().name")
    assert page.locator(".confirm-overlay input").input_value() == tab_name
    page.locator(".confirm-card .btn", has_text="Cancel").click()
    page.wait_for_selector(".confirm-overlay", state="detached")


def test_csv_download_uses_the_name_given(page):
    _run(page, "SELECT rid FROM src_1 LIMIT 3")
    page.locator(".sql-result-actions .btn", has_text="CSV").click()
    page.wait_for_selector(".confirm-overlay input")
    page.locator(".confirm-overlay input").fill("logon sweep")
    with page.expect_download() as dl:
        page.locator(".confirm-card .btn", has_text="Save").first.click()
    assert dl.value.suggested_filename == "logon sweep.csv"


def _workbook(download):
    """openpyxl refuses a path with no recognised extension, and a
    Playwright download lands in a temp file that has none."""
    import io
    from pathlib import Path as _P
    return io.BytesIO(_P(download.path()).read_bytes())


def test_excel_prompts_with_the_tab_name_like_csv_does(page):
    """Same default as CSV, for the same reason: every tab exporting to
    one filename collides in the downloads folder."""
    _run(page, "SELECT rid FROM src_1 LIMIT 3")
    page.locator(".sql-result-actions .btn", has_text="Excel").click()
    page.wait_for_selector(".confirm-overlay input")
    tab_name = page.evaluate("() => __winnow.activeSqlTab().name")
    assert page.locator(".confirm-overlay input").input_value() == tab_name
    page.locator(".confirm-card .btn", has_text="Cancel").click()
    page.wait_for_selector(".confirm-overlay", state="detached")


def test_excel_downloads_a_real_workbook_of_the_displayed_rows(page):
    """End to end: the bytes that reach the disk open as a workbook, and
    hold the header plus exactly the rows the pane was showing."""
    from openpyxl import load_workbook

    _run(page, "SELECT rid, EventId FROM src_1 ORDER BY rid LIMIT 4")
    page.locator(".sql-result-actions .btn", has_text="Excel").click()
    page.wait_for_selector(".confirm-overlay input")
    page.locator(".confirm-overlay input").fill("logon sweep")
    with page.expect_download() as dl:
        page.locator(".confirm-card .btn", has_text="Save").first.click()
    assert dl.value.suggested_filename == "logon sweep.xlsx"
    ws = load_workbook(_workbook(dl.value)).active
    assert [c.value for c in ws[1]] == ["rid", "EventId"]
    assert ws.max_row == 5, "header + the four rows on screen"


def test_excel_follows_a_click_sort(page):
    """The rows go up from the browser precisely so this holds: the sort
    exists nowhere else, so a server-side re-run would hand back a file
    in a different order than the screen."""
    from openpyxl import load_workbook

    _run(page, "SELECT rid FROM src_1 ORDER BY rid LIMIT 6")
    # Second click on the header = descending.
    page.locator("#sqlResult th", has_text="rid").click()
    page.locator("#sqlResult th", has_text="rid").click()
    page.wait_for_timeout(150)
    on_screen = [int(t) for t in page.locator("#sqlResult tr td:first-child").all_inner_texts()]
    assert on_screen == sorted(on_screen, reverse=True), "the click-sort did not take"
    page.locator(".sql-result-actions .btn", has_text="Excel").click()
    page.wait_for_selector(".confirm-overlay input")
    page.locator(".confirm-overlay input").fill("sorted")
    with page.expect_download() as dl:
        page.locator(".confirm-card .btn", has_text="Save").first.click()
    ws = load_workbook(_workbook(dl.value)).active
    in_file = [ws.cell(row=i, column=1).value for i in range(2, ws.max_row + 1)]
    assert in_file == on_screen, "the workbook must match the screen's order"


def test_save_as_table_saves_the_query_behind_the_result_not_the_editor(page):
    """The button sits on a result, but it re-ran whatever the editor held
    by then: run a query, start editing the next one, click Save on the
    result still on screen — and the table landed was the half-edited
    query's, matching neither what was shown nor what was checked. The
    request is caught and aborted here, so the shared case gains no table."""
    import json as _json

    ran = "SELECT rid FROM src_1 ORDER BY rid LIMIT 3"
    _run(page, ran)
    page.locator("#sqlText").fill("SELECT rid FROM src_1 WHERE half-edited")
    sent = []

    def catch(route):
        sent.append(_json.loads(route.request.post_data or "{}"))
        route.abort()

    page.route("**/api/sql/to_table", catch)
    try:
        page.locator(".sql-result-actions .btn", has_text="Save as table").click()
        page.wait_for_selector(".confirm-overlay input")
        page.locator(".confirm-overlay input").fill("from the result")
        page.locator(".confirm-overlay input").press("Enter")
        page.wait_for_function("() => document.querySelector('.confirm-overlay') === null")
        page.wait_for_timeout(300)
    finally:
        page.unroute("**/api/sql/to_table", catch)
    assert sent and sent[0]["sql"] == ran


def test_csv_neutralises_formulas_and_quotes_a_bare_cr_like_the_server(page):
    """The pane's CSV is built in the browser, and it wrote a formula-
    leading value (or column name) as-is — while the Excel button beside
    it, built server-side, prefixed it. A bare CR was left unquoted, which
    a spreadsheet reads as a row break. Same rules as store._csv_safe and
    Python's csv module now: strings opening with = + - @ tab or CR get a
    leading apostrophe; numbers are numbers; CR quotes like LF does."""
    from pathlib import Path as _P

    _run(page, """SELECT '=1+1' AS "=HYPERLINK(x)", 'a' || char(13) || 'b' AS cr, """
               """-5 AS n, '-5' AS s, '@x' AS a FROM src_1 LIMIT 1""")
    page.locator(".sql-result-actions .btn", has_text="CSV").click()
    page.wait_for_selector(".confirm-overlay input")
    page.locator(".confirm-overlay input").fill("formulas")
    with page.expect_download() as dl:
        page.locator(".confirm-card .btn", has_text="Save").first.click()
    text = _P(dl.value.path()).read_bytes().decode("utf-8")
    header, row = text.rstrip("\n").split("\n")
    assert header == "'=HYPERLINK(x),cr,n,s,a"
    assert row == "'=1+1,\"a\rb\",-5,'-5,'@x"
