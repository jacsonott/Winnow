"""A worksheet cannot hold every byte a case file can.

Winnow keeps control characters in cell VALUES on purpose: ingest strips
them from column names only, because a name has to be quotable while
cells are evidence and stay as they are. openpyxl refuses those same
characters, and refuses the whole workbook rather than the one cell --
so before this, a single BEL inside a single command line took down the
entire tagged export with an IllegalCharacterError naming the value but
not the row, and nothing an analyst could do about it.

They are escaped rather than dropped: these land in exactly the fields
worth reading closely, and a silently shortened string in an exported
artifact is worse than a visible \\x07.
"""

from __future__ import annotations

import io

from openpyxl import load_workbook

from winnow.store import XLSX_MAX_ROWS, _xlsx_safe, rows_to_xlsx

# BEL, SOH, VT, FF: all legal in a SQLite text column, none legal in XML.
NASTY = "cmd.exe /c echo \x07 \x01 \x0b \x0c done"


# -------------------------------------------------------- the sanitizer

def test_control_characters_are_escaped_not_dropped():
    out = _xlsx_safe(NASTY)
    assert "\\x07" in out and "\\x01" in out and "\\x0b" in out and "\\x0c" in out
    assert "\x07" not in out, "the raw character must not survive into the cell"
    # Nothing else moved: the readable text is still readable.
    assert out.startswith("cmd.exe /c echo ") and out.endswith(" done")


def test_the_three_legal_whitespace_characters_are_kept():
    """Tab, newline and carriage return are valid XML and meaningful in a
    cell. Escaping them would mangle every multi-line value in the case."""
    assert _xlsx_safe("a\tb\nc\rd") == "a\tb\nc\rd"


def test_it_still_does_what_the_csv_sanitizer_does():
    """_xlsx_safe wraps _csv_safe rather than replacing it — formula
    injection is as much a risk in .xlsx as in .csv."""
    assert _xlsx_safe("=SUM(A1)") == "'=SUM(A1)"
    assert _xlsx_safe("@cmd") == "'@cmd"


def test_a_value_with_nothing_wrong_is_returned_unchanged():
    assert _xlsx_safe("ordinary text") == "ordinary text"
    assert _xlsx_safe(None) is None
    assert _xlsx_safe(42) == 42


# ------------------------------------------------- the result workbook

def test_a_result_round_trips_through_a_workbook():
    buf = rows_to_xlsx(["Computer", "EventId"], [["DC01", 4624], ["WKSTN-02", 4625]], "Logons")
    ws = load_workbook(io.BytesIO(buf.read())).active
    assert ws.title == "Logons"
    assert [c.value for c in ws[1]] == ["Computer", "EventId"]
    assert [c.value for c in ws[2]] == ["DC01", 4624]
    assert ws.max_row == 3


def test_a_result_carrying_control_characters_still_writes():
    """The bug, at the level the analyst meets it: one bad cell used to
    cost the whole file."""
    buf = rows_to_xlsx(["CommandLine"], [[NASTY], ["clean row"]])
    ws = load_workbook(io.BytesIO(buf.read())).active
    assert "\\x07" in ws.cell(row=2, column=1).value
    assert ws.cell(row=3, column=1).value == "clean row"


def test_a_sheet_name_excel_would_refuse_is_cleaned():
    """Query tabs are named by the analyst, and / : * ? [ ] are ordinary
    things to type. The name is the default sheet name."""
    buf = rows_to_xlsx(["a"], [["x"]], "logons: DC01/DC02 [draft]?")
    ws = load_workbook(io.BytesIO(buf.read())).active
    assert ws.title == "logons_ DC01_DC02 _draft__"[:31]
    assert len(ws.title) <= 31


def test_more_rows_than_a_worksheet_holds_continue_on_the_next_sheet(monkeypatch):
    """The SQL pane no longer caps a result, so a result past Excel's
    per-sheet limit is reachable from the Excel button. Every row lands:
    the overflow continues on "Name (2)", header repeated, in order.
    Shrink the cap rather than build a million rows."""
    monkeypatch.setattr("winnow.store.XLSX_MAX_ROWS", 5)   # 4 data rows per sheet
    buf = rows_to_xlsx(["a"], [[i] for i in range(10)], "Logons")
    wb = load_workbook(io.BytesIO(buf.read()))
    assert wb.sheetnames == ["Logons", "Logons (2)", "Logons (3)"]
    got = []
    for ws in wb.worksheets:
        values = [r[0] for r in ws.iter_rows(values_only=True)]
        assert values[0] == "a", "each continuation repeats the header"
        got += values[1:]
    assert got == list(range(10)), "every row, once, in order"
    assert XLSX_MAX_ROWS == 1_048_576, "the real cap is Excel's, not the test's"


def test_an_empty_result_is_still_a_workbook_with_its_header():
    wb = load_workbook(io.BytesIO(rows_to_xlsx(["a", "b"], []).read()))
    assert wb.sheetnames == ["Query"]
    assert [c.value for c in wb.active[1]] == ["a", "b"]


# --------------------------------------------- the exports that crashed

def _tagged_case(store, tmp_path, value):
    p = tmp_path / "evt.csv"
    p.write_text(f'Line,CommandLine\n1,"{value}"\n2,"ordinary"\n', encoding="utf-8", newline="")
    rec = store.ingest_csv(str(p), name="evt.csv")
    tag = store.upsert_tag(None, "TA", "#f00", None)
    store.set_tags(rec["id"], [1, 2], tag["id"], True)
    return rec


def test_the_tagged_export_survives_a_control_character(store, tmp_path):
    _tagged_case(store, tmp_path, NASTY)
    ws = load_workbook(io.BytesIO(store.export_tagged_xlsx().read())).active
    cells = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
    assert any("\\x07" in c for c in cells), cells


def test_the_all_tables_export_survives_a_control_character(store, tmp_path):
    _tagged_case(store, tmp_path, NASTY)
    out = tmp_path / "all.xlsx"
    store.export_all_xlsx(str(out))
    ws = load_workbook(out).active
    cells = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
    assert any("\\x07" in c for c in cells), cells


def test_the_stored_value_is_never_touched(store, tmp_path):
    """Invariant #1 in spirit: the export is a copy. The case file still
    holds the bytes that were imported."""
    rec = _tagged_case(store, tmp_path, NASTY)
    store.export_tagged_xlsx()
    stored = store.run_sql(f'SELECT CommandLine FROM src_{rec["id"]} WHERE rid=1')["rows"][0][0]
    assert "\x07" in stored, "the case file must keep the raw byte"


# ------------------------------------------------------------- the route

def _held(client, sql):
    """Run in the pane and return the held result's id — the SQL pane's
    Excel… exports from there now, not from rows the browser posts up."""
    r = client.post("/api/sql/run", json={"sql": sql})
    assert r.status_code == 200, r.text
    return r.json()["id"]


def test_the_route_returns_a_workbook(client):
    rid = _held(client, "SELECT 'DC01' AS Computer, 4624 AS EventId")
    r = client.get(f"/api/sql/result/{rid}/export.xlsx?filename=logons.xlsx&sheet=Logons")
    assert r.status_code == 200
    assert r.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert 'filename="logons.xlsx"' in r.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Logons"]
    assert [c.value for c in wb.active[1]] == ["Computer", "EventId"]
    assert [c.value for c in wb.active[2]] == ["DC01", 4624], "a number stays a number"


def test_the_route_escapes_rather_than_500s(client):
    rid = _held(client, "SELECT 'cmd.exe /c echo ' || char(7) || char(1) || ' done' AS c")
    r = client.get(f"/api/sql/result/{rid}/export.xlsx")
    assert r.status_code == 200
    ws = load_workbook(io.BytesIO(r.content)).active
    assert "\\x07" in ws.cell(row=2, column=1).value


def test_the_route_spills_an_oversized_result_onto_more_sheets(client, monkeypatch):
    monkeypatch.setattr("winnow.store.XLSX_MAX_ROWS", 3)
    rid = _held(client, "SELECT 1 AS a UNION ALL SELECT 2 UNION ALL SELECT 3")
    r = client.get(f"/api/sql/result/{rid}/export.xlsx")
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Query", "Query (2)"]
    assert [c.value for c in wb["Query (2)"]["A"]] == ["a", 3]


def test_the_route_404s_a_result_that_is_gone(client):
    r = client.get("/api/sql/result/nope/export.xlsx")
    assert r.status_code == 404
