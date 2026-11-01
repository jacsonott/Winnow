"""The SQL pane's results, held server-side (winnow/sqlresults.py).

The pane has no row cap. The first build of that returned every row in
one JSON response, and the release review measured what that costs:
`SELECT *` over 2M rows of a 10-column table was 516 MB of JSON and
3.4 GB of server memory, a few more columns went past the browser's
string limit, a recursive CTE missing its stop condition grew until the
OS killed Winnow, and an Excel export of 2M rows crashed the tab. A
runaway query also could not be stopped, and kept Winnow running —
holding the case lock — after Shut down.

So a run lands its rows in a scratch table a batch at a time and the
pane pages it. These tests pin the symptoms in the terms they were
reported in: memory that grows with the answer, a response that carries
the answer, a query nothing can stop."""

from __future__ import annotations

import csv
import io
import os
import threading
import time
import tracemalloc

import pytest
from openpyxl import load_workbook

from winnow import sqlresults
from winnow.sqlresults import ResultExpired, SqlResults
from winnow.store import VIEWS_DB_SUFFIXES, VIEWS_LOCK_SUFFIX, VIEWS_PREFIX, OpCancelled, sweep_orphan_views


def _count(n, extra=""):
    return (f"WITH RECURSIVE n(i) AS (SELECT 1 UNION ALL SELECT i + 1 FROM n WHERE i < {n}) "
            f"SELECT i{extra} FROM n")


RUNAWAY = "WITH RECURSIVE n(i) AS (SELECT 1 UNION ALL SELECT i + 1 FROM n) SELECT i FROM n"
RUNAWAY_COUNT = "WITH RECURSIVE n(i) AS (SELECT 1 UNION ALL SELECT i + 1 FROM n) SELECT count(*) FROM n"


def _in_thread(fn, *args, **kwargs):
    out = {}

    def run():
        try:
            out["value"] = fn(*args, **kwargs)
        except BaseException as e:  # noqa: BLE001 — handed back to the test
            out["error"] = e
    t = threading.Thread(target=run, daemon=True)
    t.start()
    return t, out


def _wait(pred, timeout=10.0):
    end = time.time() + timeout
    while time.time() < end:
        if pred():
            return True
        time.sleep(0.02)
    return False


# ------------------------------------------------------------ the run


def test_a_run_holds_every_row_and_sends_only_the_first_page(store):
    res = store.run_sql_result(_count(12_000))
    assert res["total"] == 12_000
    assert res["columns"] == ["i"]
    assert len(res["rows"]) == SqlResults.PREVIEW_ROWS
    assert res["rows"][0] == [1] and res["rows"][-1] == [SqlResults.PREVIEW_ROWS]
    tail = store.sql_results.page(res["id"], 11_990, 500)
    assert [r[0] for r in tail["rows"]] == list(range(11_991, 12_001))


def test_memory_does_not_grow_with_the_answer(store):
    """The symptom: server memory that scaled with the result (3.4 GB for
    2M rows). 100k rows of 300-character text is ~30 MB of strings held
    whole; landed a batch at a time, the run's peak stays a few MB."""
    sql = _count(100_000, ", printf('%0300d', i) AS wide")
    tracemalloc.start()
    try:
        res = store.run_sql_result(sql)
        _, peak = tracemalloc.get_traced_memory()
    finally:
        tracemalloc.stop()
    assert res["total"] == 100_000
    assert peak < 12 * 1024 * 1024, f"peak {peak / 1e6:.1f} MB — the run is holding the answer"


def test_a_page_is_bounded_however_wide_the_rows(store, monkeypatch):
    monkeypatch.setattr(SqlResults, "PAGE_BYTES", 10_000)
    res = store.run_sql_result(_count(200, ", printf('%01000d', i) AS wide"))
    page = store.sql_results.page(res["id"], 0, 200)
    assert 1 <= len(page["rows"]) < 20, "a page past its byte budget stops early"


def test_widths_come_from_every_row(store):
    res = store.run_sql_result(_count(9_000, ", CASE WHEN i = 8999 THEN printf('%050d', i) ELSE 'x' END AS v"))
    assert res["widths"] == [4, 50]


def test_values_that_json_cannot_carry_still_page(store):
    """Spelled as /api/sql spells them (pane_cell_json): a held result and
    an inline one must not show the same cell two ways."""
    res = store.run_sql_result("SELECT CAST('abc' AS BLOB) AS b, X'FF00' AS raw, 1e999 AS big, "
                               "-1e999 AS small, 9007199254740993 AS ft")
    assert res["rows"] == [["abc", "X'FF00'", "Inf", "-Inf", "9007199254740993"]]
    assert store.sql_results.page(res["id"], 0, 1)["rows"] == res["rows"]


def test_the_pane_is_still_read_only(store):
    with pytest.raises(Exception):
        store.run_sql_result("DELETE FROM sources")
    with pytest.raises(ValueError):
        store.run_sql_result("ATTACH DATABASE 'x.db' AS x")
    assert not store.sql_results._results, "a refused statement leaves no result behind"


# ------------------------------------------------------------ stopping


def test_stop_keeps_the_rows_a_runaway_had_returned(store):
    """A recursive CTE missing its stop condition: develop returned 5,000
    rows of it instantly, the uncapped pane grew until the OS killed
    Winnow. Now it lands rows until Stop, and keeps them."""
    t, out = _in_thread(store.run_sql_result, RUNAWAY, op_token="tok-runaway")
    assert _wait(lambda: store.sql_run_progress("tok-runaway")["rows"] > 20_000)
    store.cancel_op("tok-runaway")
    t.join(10)
    assert not t.is_alive(), "Stop did not stop it"
    res = out["value"]
    assert res["stopped"] == "cancelled"
    assert res["total"] >= 20_000
    assert store.sql_results.page(res["id"], 0, 3)["rows"] == [[1], [2], [3]]
    assert store.sql_run_progress("tok-runaway") == {"running": False, "rows": 0}


def test_stop_before_the_first_row_is_a_cancel(store):
    t, out = _in_thread(store.run_sql_result, RUNAWAY_COUNT, op_token="tok-count")
    assert _wait(lambda: store.sql_run_progress("tok-count")["running"])
    time.sleep(0.2)
    store.cancel_op("tok-count")
    t.join(10)
    assert not t.is_alive()
    assert isinstance(out.get("error"), OpCancelled)


def test_shut_down_stops_a_runaway_run(store):
    """The repro: a runaway query, then Shut down. uvicorn waits for the
    request, the request never ends, Winnow stays up holding the case
    lock. Shut down's signal handler stops every pane query
    (stop_pane_queries — tests/test_sql_runaway.py drives that end to
    end); the pane's own runs have to be among them."""
    t, out = _in_thread(store.run_sql_result, RUNAWAY_COUNT, op_token="tok-shutdown")
    assert _wait(lambda: store.sql_run_progress("tok-shutdown")["running"])
    time.sleep(0.2)
    assert store.stop_pane_queries() == 1
    t.join(10)
    assert not t.is_alive(), "the query outlived Shut down"
    assert isinstance(out.get("error"), OpCancelled)


def test_closing_the_case_stops_a_run_and_deletes_the_scratch_file(case_path):
    from winnow.store import Store

    st = Store(case_path)
    res = st.run_sql_result(_count(10))
    path = st.sql_results.path
    assert path and os.path.exists(path)
    t, out = _in_thread(st.run_sql_result, RUNAWAY, op_token="tok-close")
    assert _wait(lambda: st.sql_run_progress("tok-close")["rows"] > 0)
    t0 = time.time()
    st.close()
    assert time.time() - t0 < 10
    t.join(10)
    assert not t.is_alive()
    assert not os.path.exists(path)
    with pytest.raises(ResultExpired):
        st.sql_results.page(res["id"], 0, 5)


def test_a_run_stops_short_of_filling_the_disk(store, monkeypatch):
    monkeypatch.setattr(SqlResults, "MIN_FREE_BYTES", 1 << 62)
    monkeypatch.setattr(SqlResults, "DISK_CHECK_EVERY", 1)
    res = store.run_sql_result(_count(20_000))
    assert res["stopped"] == "disk"
    assert res["total"] == SqlResults.BATCH, "what landed is kept"


# ------------------------------------------------------------ sorting


def test_the_sort_matches_the_one_the_pane_used_to_do_in_the_browser(store):
    """NULLs first, numbers by value — including numbers stored as text,
    which is every number ingest stores — then text without case, ties
    in query order. Descending flips all of it but the tie-break."""
    sql = ("SELECT v, k FROM (SELECT '10' AS v, 1 AS k UNION ALL SELECT '9', 2 UNION ALL SELECT NULL, 3 "
           "UNION ALL SELECT 'beta', 4 UNION ALL SELECT 'Alpha', 5 UNION ALL SELECT 2.5, 6 "
           "UNION ALL SELECT '9', 7 UNION ALL SELECT '', 8)")
    res = store.run_sql_result(sql)
    asc = store.sql_results.page(res["id"], 0, 10, sort=0, order="asc")["rows"]
    assert [r[1] for r in asc] == [3, 6, 2, 7, 1, 8, 5, 4]
    desc = store.sql_results.page(res["id"], 0, 10, sort=0, order="desc")["rows"]
    assert [r[1] for r in desc] == [4, 5, 8, 1, 2, 7, 6, 3]


def test_a_sorted_page_deep_in_a_big_result(store):
    res = store.run_sql_result(_count(30_000))
    page = store.sql_results.page(res["id"], 29_995, 10, sort=0, order="desc")
    assert [r[0] for r in page["rows"]] == [5, 4, 3, 2, 1]


# ------------------------------------------------------------ tags and selection


def test_tags_join_only_the_rows_the_result_holds(ingested):
    store, sid = ingested
    tag = store.list_tags()[0]
    store.set_tags(sid, [2, 4], tag["id"], True)
    res = store.run_sql_result(f"SELECT rid, User FROM src_{sid} WHERE rid <> 4")
    got = store.sql_results.tags(res["id"], {"rid_col": 0, "sid": sid}, store.path)
    assert got == {"map": {f"{sid}:2": [tag["id"]]}, "sids": [sid]}


def test_tags_join_through_a_source_id_column(ingested):
    store, sid = ingested
    tag = store.list_tags()[0]
    store.set_tags(sid, [3], tag["id"], True)
    res = store.run_sql_result(f"SELECT {sid} AS source_id, CAST(rid AS TEXT) AS rid FROM src_{sid}")
    got = store.sql_results.tags(res["id"], {"rid_col": 1, "sid_col": 0}, store.path)
    assert got["map"] == {f"{sid}:3": [tag["id"]]}
    assert got["sids"] == [sid]


def test_keys_for_a_shift_range_over_rows_never_loaded(store):
    res = store.run_sql_result(_count(20_000, ", 7 AS source_id"))
    keys = store.sql_results.keys(res["id"], 19_997, 5, {"rid_col": 0, "sid_col": 1}, sort=0, order="desc")
    assert keys == [[7, 3], [7, 2], [7, 1]]


def test_copy_of_a_selection_is_in_displayed_order_with_a_header(ingested):
    store, sid = ingested
    res = store.run_sql_result(f"SELECT rid, EventId, 1.0 AS one FROM src_{sid}")
    got = store.sql_results.pick(res["id"], sort=0, order="desc", keys=[[sid, 1], [sid, 3]],
                                 ref={"rid_col": 0, "sid": sid})
    assert got["text"].split("\n") == ["rid\tEventId\tone", "3\t4625\t1", "1\t4624\t1"]


# ------------------------------------------------------------ release


def test_a_released_result_is_gone(client, store):
    rid = client.post("/api/sql/run", json={"sql": _count(10)}).json()["id"]
    assert client.get(f"/api/sql/result/{rid}/rows?start=0&count=5").status_code == 200
    assert client.post("/api/sql/result/release", json={"ids": [rid]}).json() == {"released": 1}
    r = client.get(f"/api/sql/result/{rid}/rows?start=0&count=5")
    assert r.status_code == 404
    assert "run the query again" in r.json()["detail"]


def test_a_rerun_releases_the_result_it_replaces(client, store):
    first = client.post("/api/sql/run", json={"sql": _count(10)}).json()["id"]
    client.post("/api/sql/run", json={"sql": _count(10), "release": [first]})
    assert client.get(f"/api/sql/result/{first}/rows").status_code == 404


def test_only_the_newest_results_are_kept(store, monkeypatch):
    monkeypatch.setattr(SqlResults, "KEEP", 3)
    ids = [store.run_sql_result(_count(5))["id"] for _ in range(4)]
    with pytest.raises(ResultExpired):
        store.sql_results.page(ids[0], 0, 5)
    assert store.sql_results.page(ids[-1], 0, 5)["rows"]


# ------------------------------------------------------------ the scratch file


def test_the_scratch_file_is_swept_like_a_views_file(store):
    """Named under the views prefix and flock'd while its Store lives, so
    the startup sweep collects one a killed process left and never a live
    one."""
    assert sqlresults.FILE_PREFIX.startswith(VIEWS_PREFIX)
    store.run_sql_result(_count(3))
    path = store.sql_results.path
    assert os.path.basename(path).startswith(VIEWS_PREFIX)
    sweep_orphan_views()
    assert os.path.exists(path), "the sweep took a live Store's results"


def test_the_scratch_files_lock_is_beside_it_never_on_it(store):
    """As the views database's (test_maintenance.py): on macOS a flock on
    the database blocks SQLite's own locks on it, and the pane's first run
    failed with "database is locked"."""
    store.run_sql_result(_count(3))
    path, fd = store.sql_results.path, store.sql_results._fd
    assert fd is not None, "this filesystem took no lock"
    assert os.fstat(fd).st_ino == os.stat(path + VIEWS_LOCK_SUFFIX).st_ino != os.stat(path).st_ino
    store.sql_results.close()
    left = [s for s in VIEWS_DB_SUFFIXES + (VIEWS_LOCK_SUFFIX,) if os.path.exists(path + s)]
    assert left == [], "close() left files behind"


def test_no_scratch_file_until_the_pane_is_used(store):
    assert store.sql_results.path is None


# ------------------------------------------------------------ routes


def test_the_run_route(client, store):
    r = client.post("/api/sql/run", json={"sql": _count(7_000)})
    assert r.status_code == 200
    body = r.json()
    assert body["total"] == 7_000 and len(body["rows"]) == SqlResults.PREVIEW_ROWS
    assert body["stopped"] is None and isinstance(body["elapsed_ms"], int)


def test_the_run_route_400s_a_bad_query(client):
    r = client.post("/api/sql/run", json={"sql": "SELECT * FROM no_such_table"})
    assert r.status_code == 400
    assert "no_such_table" in r.json()["detail"]


def test_the_inline_route_is_capped_again_by_default(client):
    """/api/sql is winnow.api.sql's and the dashboards'. Uncapped by
    default, it was one JSON response of every row for anything that
    forgot to pass a limit."""
    r = client.post("/api/sql", json={"sql": _count(7_000)})
    assert len(r.json()["rows"]) == 5_000 and r.json()["truncated"] is True


def test_the_csv_export_streams_every_row_in_displayed_order(client):
    rid = client.post("/api/sql/run", json={"sql": _count(7_000, ", '=cmd' AS f, i * 1.0 AS x")}).json()["id"]
    r = client.get(f"/api/sql/result/{rid}/export.csv?sort=0&order=desc&filename=logon%20sweep.csv")
    assert r.status_code == 200
    assert "logon%20sweep.csv" in r.headers["content-disposition"]
    rows = list(csv.reader(io.StringIO(r.text)))
    assert rows[0] == ["i", "f", "x"]
    assert len(rows) == 7_001
    assert rows[1] == ["7000", "'=cmd", "7000"], "as shown, sorted, and formula-safe"


def test_the_csv_guards_column_names_but_leaves_numbers_alone(client):
    """A column named =HYPERLINK(...) was a live formula in the header
    while every cell under it was neutralised; and -5 is a number, which
    an apostrophe would have turned into text in the spreadsheet."""
    rid = client.post("/api/sql/run", json={
        "sql": """SELECT '=1+1' AS "=HYPERLINK(x)", -5 AS n, -2.5 AS r, '-5' AS s"""}).json()["id"]
    rows = list(csv.reader(io.StringIO(client.get(f"/api/sql/result/{rid}/export.csv").text)))
    assert rows == [["'=HYPERLINK(x)", "n", "r", "s"], ["'=1+1", "-5", "-2.5", "'-5"]]


def test_the_xlsx_export_spills_onto_more_sheets(client, monkeypatch):
    monkeypatch.setattr("winnow.store.XLSX_MAX_ROWS", 3)
    rid = client.post("/api/sql/run", json={"sql": _count(5)}).json()["id"]
    r = client.get(f"/api/sql/result/{rid}/export.xlsx?filename=five.xlsx&sheet=Logons")
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Logons", "Logons (2)", "Logons (3)"]
    assert [c.value for c in wb["Logons (3)"]["A"]] == ["i", 5]


def test_the_xlsx_export_follows_the_sort(client):
    rid = client.post("/api/sql/run", json={"sql": _count(6)}).json()["id"]
    r = client.get(f"/api/sql/result/{rid}/export.xlsx?sort=0&order=desc")
    ws = load_workbook(io.BytesIO(r.content)).active
    assert [c.value for c in ws["A"]] == ["i", 6, 5, 4, 3, 2, 1]


def test_a_copy_too_big_for_a_clipboard_is_a_413(client, monkeypatch):
    monkeypatch.setattr(SqlResults, "COPY_MAX_BYTES", 1_000)
    rid = client.post("/api/sql/run", json={"sql": _count(2_000)}).json()["id"]
    r = client.post(f"/api/sql/result/{rid}/pick", json={})
    assert r.status_code == 413
    assert "CSV" in r.json()["detail"]


def test_a_bad_row_reference_is_a_400(client):
    rid = client.post("/api/sql/run", json={"sql": _count(3)}).json()["id"]
    assert client.get(f"/api/sql/result/{rid}/tags?rid_col=5&sid=1").status_code == 400
    assert client.get(f"/api/sql/result/{rid}/rows?sort=9").status_code == 400


def test_progress_route(client, store):
    t, out = _in_thread(client.post, "/api/sql/run", json={"sql": RUNAWAY, "op_token": "tok-route"})
    assert _wait(lambda: client.get("/api/sql/run/progress?token=tok-route").json()["rows"] > 0)
    client.post("/api/cancel_op", json={"token": "tok-route"})
    t.join(10)
    assert out["value"].json()["stopped"] == "cancelled"
