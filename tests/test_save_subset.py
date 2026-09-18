"""Saving a view — or a pick of its rows — as a new table of the case:
Store.save_view_as_source and POST /api/view/save_as_table.

The load-bearing assertions: rows land in VIEW order with rids contiguous
from 1 (invariant #2), the columns are the parent's — derived values
included, as plain columns — and never rid/source_id/Tags/Note, a subset of
a MERGE resolves every row through its own member (invariant #9), the
subset starts untagged unless copy_tags is asked for — and that seed is
never recorded as an undo entry (invariant #7's documented exception) — a
still-building derived column refuses the copy by name, and a cancel, an
eviction or the case closing mid-copy leaves no half-table behind."""

from __future__ import annotations

import sqlite3
import threading
import time

import pytest

from winnow.store import DEFAULT_TAGS, OpCancelled, Store

ROWS = [
    ["Timestamp", "EventId", "Host"],
    ["2024-01-05 10:00:00", "4624", "h1"],
    ["2024-01-05 11:00:00", "4688", "h2"],
    ["2024-01-06 09:00:00", "4624", "h3"],
    ["2024-01-06 12:00:00", "4625", "h1"],
    ["2024-01-07 08:00:00", "4624", "h2"],
]

BY_HOST = [{"column": "Host", "dir": "asc"}]   # h1(1) h1(4) h2(2) h2(5) h3(3)


def _ingest(store, write_csv, name="events.csv", rows=ROWS):
    return store.ingest_csv(write_csv(rows, name), name=name, build_fts=False)["id"]


def _view(store, sid, **spec):
    return store.build_view(sid, {"source_id": sid, "filters": [], "sort": [], **spec})


def _save(store, view_id, name="sub", **kw):
    kw.setdefault("build_fts", False)
    return store.save_view_as_source(view_id, name, **kw)["source"]["id"]


def _hosts(store, sid):
    return [r[0] for r in store.db.execute(f"SELECT Host FROM src_{sid} ORDER BY rid")]


def _map(store, sid):
    return [tuple(r) for r in store.db.execute(
        "SELECT rid, parent_source_id, parent_rid FROM subset_rids WHERE source_id=? ORDER BY rid", (sid,))]


def _tag_id(store, name="TA"):
    return next(t for t in store.list_tags() if t["name"] == name)["id"]


# ------------------------------------------------------------ the copy

def test_whole_view_lands_in_view_order_with_contiguous_rids(store, write_csv):
    sid = _ingest(store, write_csv)
    v = _view(store, sid, sort=BY_HOST)
    assert store._views[v["view_id"]]["kind"] == "root"
    new = _save(store, v["view_id"], "by host")
    assert _hosts(store, new) == ["h1", "h1", "h2", "h2", "h3"]
    assert [r[0] for r in store.db.execute(f"SELECT rid FROM src_{new} ORDER BY rid")] == [1, 2, 3, 4, 5]
    assert _map(store, new) == [(1, sid, 1), (2, sid, 4), (3, sid, 2), (4, sid, 5), (5, sid, 3)]
    src = store.get_source(new)
    assert src["origin"] == "subset" and src["name"] == "by host" and src["row_count"] == 5
    meta = src["origin_meta"]
    assert meta["parent_source_id"] == sid and meta["parent_name"] == "events.csv"
    assert meta["parent_row_count"] == 5 and meta["selection"] == "view" and meta["rows"] == 5
    assert src["is_open"], "a saved subset opens as a tab like any new table"


def test_root_virtual_copies_the_whole_table_in_rid_order(store, write_csv):
    sid = _ingest(store, write_csv)
    v = _view(store, sid)
    assert store._views[v["view_id"]]["kind"] == "root_virtual"
    new = _save(store, v["view_id"])
    assert _hosts(store, new) == ["h1", "h2", "h3", "h1", "h2"]
    assert _map(store, new) == [(i, sid, i) for i in range(1, 6)]


def test_columns_are_the_parents_with_derived_values_as_plain_columns(store, write_csv):
    sid = _ingest(store, write_csv)
    res = store.add_derived_column(sid, "HostNum", "Host", "regex_extract", {"pattern": r"h(\d)"})
    store.wait_for_ingest_job(res["job_id"], timeout=30)
    parent = store.get_source(sid)
    v = _view(store, sid, sort=BY_HOST)
    new = _save(store, v["view_id"])
    sub = store.get_source(new)
    assert [c["name"] for c in sub["columns"]] == [c["name"] for c in parent["columns"]]
    assert not any(c.get("derived") for c in sub["columns"]), "derived VALUES, not definitions"
    # the parent's types, not a re-sample
    assert [c["type"] for c in sub["columns"]] == [c["type"] for c in parent["columns"]]
    assert [r[0] for r in store.db.execute(f"SELECT HostNum FROM src_{new} ORDER BY rid")] == ["1", "1", "2", "2", "3"]
    # never rid/source_id/Tags/Note as data columns — the SQL pane's leak
    physical = [r[1] for r in store.db.execute(f"PRAGMA table_info(src_{new})")]
    assert physical == ["rid", "Timestamp", "EventId", "Host", "HostNum"]
    assert not store.list_derived_columns(new)


def test_exclude_pairs_are_left_out(store, write_csv):
    sid = _ingest(store, write_csv)
    v = _view(store, sid)
    new = _save(store, v["view_id"], exclude=[[sid, 2], [sid, 4], [sid, 999]])
    assert _hosts(store, new) == ["h1", "h3", "h2"]
    assert _map(store, new) == [(1, sid, 1), (2, sid, 3), (3, sid, 5)]
    assert store.get_source(new)["origin_meta"]["excluded"] == 3


def test_explicit_keys_land_in_view_order_and_strangers_are_skipped(store, write_csv):
    sid = _ingest(store, write_csv)
    v = _view(store, sid, sort=[{"column": "Timestamp", "dir": "desc"}])
    new = _save(store, v["view_id"], keys=[[sid, 1], [sid, 5], [sid, 5], [sid, 99], [sid + 7, 1]])
    assert _map(store, new) == [(1, sid, 5), (2, sid, 1)], "view order, deduped, unknown rows dropped"
    assert store.get_source(new)["origin_meta"]["selection"] == "picks"
    # the same picks against the virtual root come back in rid order
    v2 = _view(store, sid)
    new2 = _save(store, v2["view_id"], "picks2", keys=[[sid, 5], [sid, 1]])
    assert _map(store, new2) == [(1, sid, 1), (2, sid, 5)]


def test_too_many_picks_is_a_value_error(store, write_csv):
    sid = _ingest(store, write_csv)
    v = _view(store, sid)
    with pytest.raises(ValueError):
        store.save_view_as_source(v["view_id"], "x", keys=[[sid, i] for i in range(store.SELECTION_REMAP_MAX + 1)])
    with pytest.raises(ValueError):
        store.save_view_as_source(v["view_id"], "   ")


def test_group_views_copy_just_the_group(store, write_csv):
    sid = _ingest(store, write_csv)
    root = _view(store, sid)
    g = store.expand_group(root["view_id"], "EventId", "4624")
    assert store._views[g["view_id"]]["kind"] == "group_virtual"
    new = _save(store, g["view_id"], "logons")
    assert _map(store, new) == [(1, sid, 1), (2, sid, 3), (3, sid, 5)]
    assert {r[0] for r in store.db.execute(f"SELECT EventId FROM src_{new}")} == {"4624"}
    # picks against a virtual group: only rows the group's predicate admits
    picked = _save(store, g["view_id"], "picked", keys=[[sid, 5], [sid, 2]])
    assert _map(store, picked) == [(1, sid, 5)]
    # the materialised branch (kind 'group') reads through its own v.view_N
    store.GROUP_MATERIALIZE_THRESHOLD = 0
    g2 = store.expand_group(root["view_id"], "EventId", "4624")
    assert store._views[g2["view_id"]]["kind"] == "group"
    new2 = _save(store, g2["view_id"], "logons2")
    assert _map(store, new2) == [(1, sid, 1), (2, sid, 3), (3, sid, 5)]


def test_soft_cap_asks_then_obeys_force(store, write_csv, monkeypatch):
    sid = _ingest(store, write_csv)
    v = _view(store, sid)
    monkeypatch.setattr(type(store), "SQL_TO_TABLE_SOFT_CAP", 2)
    assert store.save_view_as_source(v["view_id"], "big") == {"needs_confirm": True, "rows": 5}
    assert store.save_view_as_source(v["view_id"], "big", exclude=[[sid, 1]]) == {"needs_confirm": True, "rows": 4}
    assert store.save_view_as_source(v["view_id"], "small", keys=[[sid, 1], [sid, 2]], build_fts=False)["source"]["row_count"] == 2
    assert store.save_view_as_source(v["view_id"], "big", force=True, build_fts=False)["source"]["row_count"] == 5


def test_copy_streams_batch_by_batch(store, write_csv, monkeypatch):
    """Keyset paging across batches, not one list of the whole view — the
    map and the row count are right at every commit."""
    import winnow.store as st
    monkeypatch.setattr(st, "BATCH", 2)
    commits = []
    real = Store._commit_ingest_batch

    def counting(self, insert_sql, batch, source_id, total, **kw):
        commits.append(len(batch))
        return real(self, insert_sql, batch, source_id, total, **kw)
    monkeypatch.setattr(Store, "_commit_ingest_batch", counting)
    sid = _ingest(store, write_csv)
    for spec in ({}, {"sort": BY_HOST}):
        commits.clear()
        v = _view(store, sid, **spec)
        new = _save(store, v["view_id"], "batched" + str(len(spec)))
        assert commits == [2, 2, 1]
        assert store.get_source(new)["row_count"] == 5
        assert [m[0] for m in _map(store, new)] == [1, 2, 3, 4, 5]


# --------------------------------------------------- tags, notes, layout

def test_tags_and_notes_are_seeded_on_request_once_and_never_undoable(store, write_csv):
    """copy_tags=True is the explicit, scripts-only seed; the default is a
    clean slate (test_the_default_is_untagged_because_the_timeline_and_export_would_double_count)."""
    sid = _ingest(store, write_csv)
    ta = _tag_id(store)
    store.set_tags(sid, [1, 3], ta, True)
    store.set_note(sid, 3, "pivot host")
    store.save_layout(sid, {"columns": {"Host": {"w": 240}}, "order": ["Host", "EventId", "Timestamp"]})
    v = _view(store, sid, sort=BY_HOST)
    undo_before = store.undo_peek()
    new = _save(store, v["view_id"], copy_tags=True)
    # parent rids 1 and 3 sit at subset rids 1 and 5 under the Host sort
    assert {r[0] for r in store.db.execute("SELECT rid FROM row_tags WHERE source_id=?", (new,))} == {1, 5}
    assert store.db.execute("SELECT note FROM row_notes WHERE source_id=? AND rid=5", (new,)).fetchone()[0] == "pivot host"
    assert store.get_source(new)["tagged_row_count"] == 2 and store.get_source(new)["note_count"] == 1
    assert store.undo_peek() == undo_before, "the seed is not a tag change anyone can undo"
    assert store.get_layout(new) == store.get_layout(sid)
    # the parent's tags are untouched, and untagging the subset later stays on the subset
    store.set_tags(new, [1], ta, False)
    assert store.get_source(sid)["tagged_row_count"] == 2
    # the layout still comes along without the seed
    bare = _save(store, v["view_id"], "bare")
    assert store.get_source(bare)["tagged_row_count"] == 0 and store.get_source(bare)["note_count"] == 0
    assert store.get_layout(bare) == store.get_layout(sid)


def test_the_default_is_untagged_because_the_timeline_and_export_would_double_count(client, store, write_csv):
    """build_timeline and export_tagged_xlsx list every source's tagged
    rows — a subset's included, on purpose, since a tag put on a subset
    is real work — so a seeded copy shows each of the parent's findings
    twice. Hence the default at the store and the route: a clean slate."""
    from openpyxl import load_workbook
    sid = _ingest(store, write_csv)
    store.set_tags(sid, [1, 3], _tag_id(store), True)
    store.set_note(sid, 1, "n")
    v = _view(store, sid)
    r = client.post("/api/view/save_as_table", json={"view_id": v["view_id"], "name": "clean"})
    assert r.status_code == 200, r.text
    clean = store.get_source(r.json()["source"]["id"])
    assert clean["tagged_row_count"] == 0 and clean["note_count"] == 0
    assert store.get_source(sid)["tagged_row_count"] == 2, "the parent keeps its own"
    assert store.build_timeline()["row_count"] == 2, "each finding once"
    assert load_workbook(store.export_tagged_xlsx()).sheetnames == ["events.csv"]
    # the opt-in, and what it does to the Timeline and the whole-case
    # export: each finding twice, a second sheet of the same rows
    r = client.post("/api/view/save_as_table", json={"view_id": v["view_id"], "name": "seeded", "copy_tags": True})
    assert store.get_source(r.json()["source"]["id"])["tagged_row_count"] == 2
    assert store.build_timeline()["row_count"] == 4
    assert load_workbook(store.export_tagged_xlsx()).sheetnames == ["events.csv", "seeded"]


def test_subset_rids_go_with_the_table_in_both_directions(store, write_csv):
    sid = _ingest(store, write_csv)
    v = _view(store, sid)
    a = _save(store, v["view_id"], "a")
    b = _save(store, v["view_id"], "b")
    store.drop_source(a)
    assert _map(store, a) == [] and len(_map(store, b)) == 5
    assert store.db.execute("SELECT COUNT(*) FROM subset_rids WHERE source_id=?", (a,)).fetchone()[0] == 0
    # dropping the PARENT clears every map that named it — its id gets reused
    store.drop_source(sid)
    assert _map(store, b) == []
    assert store.get_source(b)["row_count"] == 5, "the subset itself is a real table and survives its parent"
    assert store.get_source(b)["origin_meta"]["parent_name"] == "events.csv"


# ------------------------------------------------------------ merges

def test_a_subset_of_a_merge_resolves_each_row_through_its_member(store, write_csv):
    a = _ingest(store, write_csv, "a.csv", [["When", "Msg"], ["2024-01-01 10:00", "alpha"], ["2024-01-01 12:00", "beta"]])
    b = _ingest(store, write_csv, "b.csv", [["When", "Msg"], ["2024-01-01 11:00", "gamma"], ["2024-01-01 13:00", "delta"]])
    ta = _tag_id(store)
    store.set_tags(b, [1], ta, True)
    store.set_note(a, 2, "from a")
    mid = store.create_merge("both", [a, b])["id"]
    v = _view(store, mid, sort=[{"column": "When", "dir": "asc"}])
    new = _save(store, v["view_id"], "merged slice", copy_tags=True)
    assert [r[0] for r in store.db.execute(f"SELECT Msg FROM src_{new} ORDER BY rid")] == ["alpha", "gamma", "beta", "delta"]
    assert _map(store, new) == [(1, a, 1), (2, b, 1), (3, a, 2), (4, b, 2)]
    src = store.get_source(new)
    assert [c["name"] for c in src["columns"]] == ["When", "Msg"]
    assert src["origin_meta"]["parent_source_id"] == mid and src["origin_meta"]["parent_name"] == "both"
    assert {r[0] for r in store.db.execute("SELECT rid FROM row_tags WHERE source_id=?", (new,))} == {2}
    assert store.db.execute("SELECT note FROM row_notes WHERE source_id=? AND rid=3", (new,)).fetchone()[0] == "from a"
    # explicit picks across members keep the merge's order too
    picked = _save(store, v["view_id"], "picked", keys=[[b, 2], [a, 1]])
    assert _map(store, picked) == [(1, a, 1), (2, b, 2)]


# ------------------------------------------------- failure drops the partial

def test_cancel_drops_the_partial_source(store, write_csv):
    sid = _ingest(store, write_csv)
    v = _view(store, sid)
    before = {s["id"] for s in store.list_sources()}
    store.cancel_op("subset-1")   # lands before the op registers — still takes effect
    with pytest.raises(OpCancelled):
        store.save_view_as_source(v["view_id"], "doomed", op_token="subset-1", build_fts=False)
    assert {s["id"] for s in store.list_sources()} == before
    assert store.db.execute("SELECT COUNT(*) FROM subset_rids").fetchone()[0] == 0


def test_eviction_mid_copy_is_409_shaped_and_drops_the_partial_source(store, write_csv, monkeypatch):
    sid = _ingest(store, write_csv)
    v = _view(store, sid, sort=BY_HOST)
    real = Store._subset_cells

    def evicting(self, ro, cols, chunk, op_token):
        out = real(self, ro, cols, chunk, op_token)
        store.close_view(v["view_id"])   # a filter change on the parent, mid-copy
        return out
    monkeypatch.setattr(Store, "_subset_cells", evicting)
    before = {s["id"] for s in store.list_sources()}
    with pytest.raises(KeyError, match="expired"):
        store.save_view_as_source(v["view_id"], "gone", build_fts=False)
    assert {s["id"] for s in store.list_sources()} == before


def test_a_derived_column_still_building_refuses_the_copy_by_name(client, store, write_csv):
    """Derived values are copied as they stand, so a backfill mid-way
    would land blanks in a plain column nothing can re-derive. The status
    is pinned rather than raced against the job: the rule reads
    derived_status, and that is what is set here."""
    sid = _ingest(store, write_csv)
    res = store.add_derived_column(sid, "HostNum", "Host", "regex_extract", {"pattern": r"h(\d)"})
    store.wait_for_ingest_job(res["job_id"], timeout=30)
    v = _view(store, sid)
    with store.lock, store.db:
        store.db.execute("UPDATE derived_columns SET status='building' WHERE source_id=? AND name='HostNum'", (sid,))
    before = {s["id"] for s in store.list_sources()}
    with pytest.raises(ValueError, match="HostNum"):
        store.save_view_as_source(v["view_id"], "early", build_fts=False)
    assert {s["id"] for s in store.list_sources()} == before, "refused before any shell exists"
    r = client.post("/api/view/save_as_table", json={"view_id": v["view_id"], "name": "early"})
    assert r.status_code == 400 and "HostNum" in r.json()["detail"]
    with store.lock, store.db:
        store.db.execute("UPDATE derived_columns SET status='ready' WHERE source_id=?", (sid,))
    new = _save(store, v["view_id"], "later")
    assert [r[0] for r in store.db.execute(f"SELECT HostNum FROM src_{new} ORDER BY rid")] == ["1", "2", "3", "1", "2"]


def test_a_merge_refuses_while_any_member_copy_of_a_derived_column_is_building(store, write_csv):
    """Invariant #9: a merge exposes a derived column only when every
    member has it, and each member's status is its own — one member
    still building is enough to refuse."""
    a = _ingest(store, write_csv, "a.csv")
    b = _ingest(store, write_csv, "b.csv", rows=[ROWS[0], ["2024-02-01 10:00:00", "4624", "h4"], ["2024-02-01 11:00:00", "4625", "h5"]])
    for m in (a, b):
        res = store.add_derived_column(m, "HostNum", "Host", "regex_extract", {"pattern": r"h(\d)"})
        store.wait_for_ingest_job(res["job_id"], timeout=30)
    mid = store.create_merge("both", [a, b])["id"]
    assert "HostNum" in [c["name"] for c in store._source_lite(mid)["columns"]], "canonical: on every member"
    v = _view(store, mid, sort=BY_HOST)
    with store.lock, store.db:
        store.db.execute("UPDATE derived_columns SET status='building' WHERE source_id=? AND name='HostNum'", (b,))
    with pytest.raises(ValueError, match="HostNum"):
        store.save_view_as_source(v["view_id"], "early", build_fts=False)
    with store.lock, store.db:
        store.db.execute("UPDATE derived_columns SET status='ready' WHERE source_id=?", (b,))
    new = _save(store, v["view_id"], "later")
    assert store.get_source(new)["row_count"] == 7
    assert sorted(r[0] for r in store.db.execute(f"SELECT HostNum FROM src_{new}")) == ["1", "1", "2", "2", "3", "4", "5"]


def test_close_with_a_copy_in_flight_drops_the_partial_source(case_path, write_csv, monkeypatch):
    """A case switch or a shutdown mid-copy: close() flags the copy, waits
    for it to drop its partial source, and only then closes the writer —
    the contract tests/test_ingest_jobs.py's close case pins for a job,
    for a copy that runs on the request's own thread instead."""
    import winnow.store as st
    monkeypatch.setattr(st, "BATCH", 2)
    s = Store(case_path, default_tags=DEFAULT_TAGS)
    sid = s.ingest_csv(write_csv(ROWS, "events.csv"), name="events.csv", build_fts=False)["id"]
    v = _view(s, sid, sort=BY_HOST)
    real = Store._subset_cells
    started = threading.Event()

    def stalling(self, ro, cols, chunk, op_token):
        # The first batch is read, then held until close() has flagged the
        # copy — so the cancel lands mid-copy, after one committed batch,
        # every time rather than when the scheduler feels like it.
        out = real(self, ro, cols, chunk, op_token)
        started.set()
        deadline = time.time() + 10
        while time.time() < deadline:
            with self._ingest_jobs_lock:
                if any(c["cancelled"] for c in self._subset_copies.values()):
                    break
            time.sleep(0.01)
        return out
    monkeypatch.setattr(Store, "_subset_cells", stalling)
    outcome = {}

    def run():
        try:
            s.save_view_as_source(v["view_id"], "doomed", build_fts=False)
            outcome["result"] = "finished"
        except BaseException as e:  # noqa: BLE001 — recorded for the assertion below
            outcome["error"] = e
    t = threading.Thread(target=run, daemon=True)
    t.start()
    try:
        assert started.wait(10)
        assert s.copies_in_flight() == 1, "what server.py's _busy_reason reads"
        t0 = time.time()
        s.close()  # must cancel the copy and return, not hang or crash
        assert time.time() - t0 < 30
        t.join(10)
        assert not t.is_alive()
        assert isinstance(outcome.get("error"), OpCancelled), outcome
    finally:
        if not s.closed:
            s.close()
    con = sqlite3.connect(case_path)
    try:
        assert [r[0] for r in con.execute("SELECT name FROM sources ORDER BY id")] == ["events.csv"]
        assert [r[0] for r in con.execute(
            "SELECT name FROM sqlite_master WHERE type='table' AND name LIKE 'src_%' ORDER BY name")] == [f"src_{sid}"]
        assert con.execute("SELECT COUNT(*) FROM subset_rids").fetchone()[0] == 0
        assert con.execute("SELECT COUNT(*) FROM open_tabs WHERE source_id<>?", (sid,)).fetchone()[0] == 0
    finally:
        con.close()


def test_a_copy_arriving_while_the_case_closes_backs_out_with_nothing_created(case_path, write_csv, monkeypatch):
    """The other side of the race: close() has already taken its snapshot
    when the copy creates its shell — the copy sees _closing under the
    same lock hold, drops the shell itself and gives up, rather than
    running its batches into a writer that is about to close."""
    s = Store(case_path, default_tags=DEFAULT_TAGS)
    sid = s.ingest_csv(write_csv(ROWS, "events.csv"), name="events.csv", build_fts=False)["id"]
    v = _view(s, sid)
    with s._ingest_jobs_lock:
        s._closing = True   # what close() sets in its flag-and-snapshot step
    try:
        with pytest.raises(OpCancelled):
            s.save_view_as_source(v["view_id"], "late", build_fts=False)
        assert [x["name"] for x in s.list_sources()] == ["events.csv"]
        assert s.copies_in_flight() == 0
    finally:
        s.close()


# --------------------------------------------------- the route, sessions

def test_route_round_trip_and_sources_listing_carries_provenance(client, store, write_csv):
    sid = _ingest(store, write_csv)
    v = _view(store, sid)
    r = client.post("/api/view/save_as_table", json={
        "view_id": v["view_id"], "name": "via route", "exclude": [[sid, 1]], "spec": {"filters": [1]}})
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["source"]["row_count"] == 4 and body["source"]["name"] == "via route"
    me = next(s for s in client.get("/api/sources").json() if s["id"] == body["source"]["id"])
    assert me["origin"] == "subset"
    assert me["origin_meta"]["parent_source_id"] == sid and me["origin_meta"]["spec"] == {"filters": [1]}
    assert client.post("/api/view/save_as_table", json={"view_id": v["view_id"], "name": " "}).status_code == 400


def test_an_expired_view_is_a_409_at_the_route(client, store, write_csv):
    sid = _ingest(store, write_csv)
    v = _view(store, sid, sort=BY_HOST)
    store.close_view(v["view_id"])
    r = client.post("/api/view/save_as_table", json={"view_id": v["view_id"], "name": "late"})
    assert r.status_code == 409 and "expired" in r.json()["detail"].lower()


def test_a_saved_session_reattaches_tags_on_a_subset(store, write_csv):
    sid = _ingest(store, write_csv)
    v = _view(store, sid)
    new = _save(store, v["view_id"], "kept")
    assert store.get_source(new)["file_hash"].startswith("subset:")
    ta = _tag_id(store)
    store.set_tags(new, [2, 4], ta, True)
    store.save_session("pass one")
    store.set_tags(new, [2, 4], ta, False)
    assert store.get_source(new)["tagged_row_count"] == 0
    res = store.load_session("pass one")
    assert not [w for w in res["warnings"] if "kept" in w], res["warnings"]
    assert {r[0] for r in store.db.execute("SELECT rid FROM row_tags WHERE source_id=?", (new,))} == {2, 4}
    # two subsets of the same view are two sources, not one hash
    other = _save(store, v["view_id"], "kept again")
    assert store.get_source(other)["file_hash"] != store.get_source(new)["file_hash"]


# ----------------------------------------- copy_sources_to and dashboards

def test_copy_sources_to_carries_origin_and_remaps_the_map(store, write_csv, tmp_path):
    sid = _ingest(store, write_csv)
    v = _view(store, sid, sort=BY_HOST)
    new = _save(store, v["view_id"], "slice")

    t = Store(str(tmp_path / "with-parent.db"))
    tp = t.path
    t.close()
    res = store.copy_sources_to(tp, [new, sid])   # child before parent on purpose
    ids = {c["name"]: c["id"] for c in res["copied"]}
    t = Store(tp)
    try:
        s = t.get_source(ids["slice"])
        assert s["origin"] == "subset"
        assert s["origin_meta"]["parent_source_id"] == ids["events.csv"]
        assert s["origin_meta"]["parent_name"] == "events.csv"
        assert s["file_hash"] == store.get_source(new)["file_hash"]
        rows = [tuple(r) for r in t.db.execute(
            "SELECT rid, parent_source_id, parent_rid FROM subset_rids WHERE source_id=? ORDER BY rid", (ids["slice"],))]
        assert rows == [(1, ids["events.csv"], 1), (2, ids["events.csv"], 4), (3, ids["events.csv"], 2),
                        (4, ids["events.csv"], 5), (5, ids["events.csv"], 3)]
    finally:
        t.close()

    t = Store(str(tmp_path / "alone.db"))
    tp = t.path
    t.close()
    res = store.copy_sources_to(tp, [new])
    t = Store(tp)
    try:
        s = t.get_source(res["copied"][0]["id"])
        assert s["origin"] == "subset" and s["origin_meta"]["parent_name"] == "events.csv"
        assert s["origin_meta"]["parent_source_id"] is None, "the parent stayed behind — no id to point at"
        assert t.db.execute("SELECT COUNT(*) FROM subset_rids").fetchone()[0] == 0
        assert s["row_count"] == 5
    finally:
        t.close()


def test_header_set_union_skips_subsets(store, write_csv, monkeypatch):
    from winnow import defaults
    monkeypatch.setattr(defaults, "headers", lambda: {
        "version": 1, "nicknames": [("Events", ["Timestamp", "EventId", "Host"])]})
    sid = _ingest(store, write_csv)
    v = _view(store, sid)
    _save(store, v["view_id"], "twin")
    assert [m["id"] for m in store._sources_for_header_set("Events")] == [sid]


def test_sql_to_table_records_its_query_as_origin(store, write_csv):
    sid = _ingest(store, write_csv)
    res = store.sql_to_table(f"SELECT Host FROM src_{sid} WHERE EventId = '4624'", "logon hosts")
    src = store.get_source(res["source"]["id"])
    assert src["origin"] == "sql" and src["origin_meta"] == {"sql": f"SELECT Host FROM src_{sid} WHERE EventId = '4624'"}
    assert src["file_hash"] is None
    imported = store.get_source(sid)
    assert imported["origin"] is None and imported["origin_meta"] is None


def test_an_older_case_file_gains_the_columns_on_open(tmp_path):
    """CREATE TABLE IF NOT EXISTS can't add a column: a case from before
    provenance existed is patched in place, and its rows read as imports."""
    import sqlite3
    path = str(tmp_path / "old.db")
    s = Store(path)
    s.close()
    con = sqlite3.connect(path)
    con.executescript("""
        CREATE TABLE sources_old AS SELECT id, name, path, table_name, row_count, columns,
            file_hash, imported_at, has_fts, nickname FROM sources;
        DROP TABLE sources; ALTER TABLE sources_old RENAME TO sources;
        INSERT INTO sources(id, name, path, table_name, row_count, columns, file_hash, imported_at, has_fts)
            VALUES (1, 'legacy.csv', NULL, 'src_1', 1, '[{"name":"A","type":"text"}]', NULL, '2024', 0);
        CREATE TABLE src_1 (rid INTEGER PRIMARY KEY, A TEXT); INSERT INTO src_1(A) VALUES ('x');
        DROP TABLE subset_rids;""")
    con.commit()
    con.close()
    s = Store(path)
    try:
        cols = {r[1] for r in s.db.execute("PRAGMA table_info(sources)")}
        assert {"origin", "origin_meta"} <= cols
        src = s.get_source(1)
        assert src["origin"] is None and src["origin_meta"] is None
        v = s.build_view(1, {"source_id": 1, "filters": [], "sort": []})
        assert s.save_view_as_source(v["view_id"], "from legacy", build_fts=False)["source"]["row_count"] == 1
    finally:
        s.close()
