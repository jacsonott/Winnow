"""Excel workbook ingest, end to end: the xlsxread text conversions, the
sheet preview/ingest pair on Store, the multi-sheet background job, and
the HTTP routing. Fixtures are written with openpyxl — already a runtime
dependency for the XLSX export, and the same library the reader wraps."""

from __future__ import annotations

import datetime

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

from winnow import xlsxread
from winnow.store import BATCH


@pytest.fixture
def write_xlsx(tmp_path):
    """write_xlsx({sheet: rows}) -> path. Cell values keep their Python
    types (str/int/float/bool/datetime/None) — openpyxl styles datetimes
    as dates on its own, which is exactly the signal the reader keys on."""

    def _write(sheets: dict[str, list[list]], name: str = "book.xlsx") -> str:
        wb = Workbook()
        wb.remove(wb.active)
        for title, rows in sheets.items():
            ws = wb.create_sheet(title)
            for r in rows:
                ws.append(r)
        path = tmp_path / name
        wb.save(path)
        return str(path)

    return _write


def test_cell_text_conversions():
    assert xlsxread.cell_text(None) == ""
    assert xlsxread.cell_text("plain") == "plain"
    assert xlsxread.cell_text(True) == "TRUE"
    assert xlsxread.cell_text(False) == "FALSE"
    assert xlsxread.cell_text(42) == "42"
    assert xlsxread.cell_text(42.0) == "42"
    assert xlsxread.cell_text(3.25) == "3.25"
    assert xlsxread.cell_text(datetime.datetime(2024, 5, 1, 13, 7, 9)) == "2024-05-01 13:07:09"
    # Day-serials are floats: 999999µs of storage imprecision must round,
    # not print — and exact midnight renders as a bare date.
    assert xlsxread.cell_text(datetime.datetime(2024, 5, 1, 13, 6, 59, 999999)) == "2024-05-01 13:07:00"
    assert xlsxread.cell_text(datetime.datetime(2024, 5, 1)) == "2024-05-01"
    assert xlsxread.cell_text(datetime.date(2024, 5, 2)) == "2024-05-02"
    assert xlsxread.cell_text(datetime.time(9, 30, 5)) == "09:30:05"
    assert xlsxread.cell_text(datetime.timedelta(hours=26, minutes=3, seconds=4)) == "26:03:04"
    # On the hour, and past a hundred: the hours are whole hours, however many.
    assert xlsxread.cell_text(datetime.timedelta(hours=1)) == "1:00:00"
    assert xlsxread.cell_text(datetime.timedelta(minutes=59, seconds=59)) == "0:59:59"
    assert xlsxread.cell_text(datetime.timedelta(hours=100)) == "100:00:00"


def test_reader_rows_dates_and_gaps(write_xlsx):
    path = write_xlsx({"Events": [
        ["When", "Host", "Count"],
        [datetime.datetime(2024, 5, 1, 13, 7), "srv01", 42],
        [datetime.date(2024, 5, 2), None, 3.5],  # gap stays a gap, in place
    ]})
    hint_rows, hint_cols, rows = xlsxread.sheet_reader(path, "Events")
    assert (hint_rows, hint_cols) == (3, 3)
    assert list(rows) == [
        ["When", "Host", "Count"],
        ["2024-05-01 13:07:00", "srv01", "42"],
        ["2024-05-02", "", "3.5"],
    ]
    with pytest.raises(ValueError, match="No such sheet"):
        xlsxread.sheet_reader(path, "Nope")


def test_workbook_sheets_listing(write_xlsx):
    path = write_xlsx({
        "Events": [["When", "Host"], ["a", "b"], ["c", "d"]],
        "Empty": [],
        "HeaderOnly": [["X"]],
    })
    sheets = xlsxread.workbook_sheets(path)
    assert [s["name"] for s in sheets] == ["Events", "Empty", "HeaderOnly"]
    assert sheets[0] == {"name": "Events", "row_count": 2, "columns": ["When", "Host"]}
    assert sheets[1]["row_count"] == 0 and sheets[1]["columns"] == []
    assert sheets[2]["row_count"] == 0 and sheets[2]["columns"] == ["X"]


def test_preview_matches_sqlite_picker_shape(store, write_xlsx):
    path = write_xlsx({"Log": [["Time", "Msg"], ["t1", "m1"]]})
    res = store.preview_xlsx_sheets(path)
    assert res == {"tables": [{
        "name": "Log", "row_count": 1,
        "columns": [{"name": "Time", "type": "TEXT"}, {"name": "Msg", "type": "TEXT"}],
        "likely_timestamp_columns": [],
    }]}


def test_ingest_sheet_values_and_naming(store, write_xlsx):
    path = write_xlsx({"Events": [
        [None, None],                       # leading empty row — skipped
        ["When", "Host", "Host"],           # dupe headers go through sanitize_columns
        [datetime.datetime(2024, 5, 1, 8, 0, 30), "srv01", "x"],
        ["not-a-date", None, "y"],
    ]}, name="triage.xlsx")
    rec = store.ingest_xlsx_sheet(path, "Events", build_fts=False)
    assert rec["name"] == "triage.Events"
    assert rec["row_count"] == 2
    assert rec["ragged_rows"] == 0
    assert [c["name"] for c in rec["columns"]] == ["When", "Host", "Host_1"]
    got = store.db.execute(f"SELECT * FROM src_{rec['id']} ORDER BY rid").fetchall()
    assert [tuple(r)[1:] for r in got] == [
        ("2024-05-01 08:00:30", "srv01", "x"),
        ("not-a-date", "", "y"),
    ]


def test_ingest_empty_sheet_refused(store, write_xlsx):
    path = write_xlsx({"Blank": []})
    with pytest.raises(ValueError, match="has no rows"):
        store.ingest_xlsx_sheet(path, "Blank")
    assert store.list_sources() == []


def test_multi_sheet_job(store, write_xlsx):
    path = write_xlsx({
        "A": [["X"], ["1"], ["2"]],
        "B": [["Y"], ["3"]],
    })
    job = store.start_ingest_job("xlsx", path, options={
        "build_fts": False,
        "tables": [{"table": "A"}, {"table": "B", "name": "renamed"}],
    })
    assert job["tables_total"] == 2
    done = store.wait_for_ingest_job(job["job_id"], timeout=30)
    assert done["status"] == "done"
    assert done["tables_done"] == 2
    assert done["rows_done"] == 3
    names = [s["name"] for s in store.list_sources()]
    assert "book.A" in names and "renamed" in names


def test_cancel_drops_partial_sheet(store, write_xlsx):
    # More rows than one BATCH so the per-batch cancel hook actually fires.
    rows = [["N"]] + [[i] for i in range(BATCH + 50)]
    path = write_xlsx({"Big": rows})
    from winnow.store import IngestCancelled
    with pytest.raises(IngestCancelled):
        store.ingest_xlsx_sheet(path, "Big", build_fts=False, cancel=lambda: True)
    assert store.list_sources() == []


def test_http_routing(store, write_xlsx, monkeypatch):
    import server
    monkeypatch.setattr(server, "STORE", store)
    client = TestClient(server.app)
    headers = {"X-Timeline-Lite-Client": "1"}
    path = write_xlsx({"Log": [["A"], ["v"]]})

    # Extension routing: .xlsx auto-detects, preview returns the sheet list.
    res = client.post("/api/ingest/preview/path", json={"path": path}, headers=headers)
    assert res.status_code == 200
    assert [t["name"] for t in res.json()["tables"]] == ["Log"]

    # The job route refuses a sheetless xlsx import, same gate as sqlite.
    res = client.post("/api/ingest/jobs/path", json={"path": path}, headers=headers)
    assert res.status_code == 400
    assert "sheets" in res.json()["detail"]

    # Upload preview spools the whole workbook (a zip can't be prefix-sniffed).
    with open(path, "rb") as f:
        res = client.post("/api/ingest/xlsx/preview", headers=headers,
                          files={"file": ("book.xlsx", f)})
    assert res.status_code == 200
    assert res.json()["tables"][0]["row_count"] == 1

    res = client.post("/api/ingest/jobs/path",
                      json={"path": path, "tables": [{"table": "Log"}], "build_fts": False},
                      headers=headers)
    assert res.status_code == 200
    done = store.wait_for_ingest_job(res.json()["job_id"], timeout=30)
    assert done["status"] == "done" and done["rows_done"] == 1

    # And the upload transport: spool + background job, same tables option.
    with open(path, "rb") as f:
        res = client.post("/api/ingest/jobs/upload", headers=headers,
                          files={"file": ("book.xlsx", f)},
                          data={"kind": "xlsx", "build_fts": "false",
                                "tables": '[{"table": "Log", "name": "uploaded"}]'})
    assert res.status_code == 200
    done = store.wait_for_ingest_job(res.json()["job_id"], timeout=30)
    assert done["status"] == "done"
    assert "uploaded" in [s["name"] for s in store.list_sources()]


def test_an_out_of_range_date_cell_keeps_its_number(tmp_path):
    """openpyxl turns a date-styled cell whose serial no datetime can hold
    (year 10000 and on) into the string "#VALUE!", and the number
    is gone — a value the workbook stores imported as an error. The reader
    does the date conversion itself now and keeps the number when the
    date can't be made; real dates and [h]:mm durations convert as before."""
    import warnings

    wb = Workbook()
    ws = wb.active
    ws.append(["when", "elapsed", "n"])
    ws.append([datetime.datetime(2024, 5, 1, 13, 7, 0), datetime.timedelta(hours=26, minutes=5), 1])
    ws.cell(row=3, column=1, value=99999999).number_format = "yyyy-mm-dd"
    ws.cell(row=4, column=1, value=2958466).number_format = "yyyy-mm-dd hh:mm"   # 10000-01-01
    ws.cell(row=3, column=3, value=7)
    path = tmp_path / "range.xlsx"
    wb.save(path)

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        rows = list(xlsxread.iter_sheet_rows(str(path), ws.title))
    assert rows[1] == ["2024-05-01 13:07:00", "26:05:00", "1"]
    assert rows[2][0] == "99999999" and rows[2][2] == "7"
    assert rows[3][0] == "2958466"
    assert "#VALUE!" not in [c for r in rows for c in r]


def _with_declared_dimension(path, ref, out):
    """A copy of the workbook at `path` whose first sheet declares
    <dimension ref=`ref`> — stale or defaulted, as several writers emit."""
    import re
    import zipfile

    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = re.sub(rb'<dimension ref="[^"]*"', b'<dimension ref="' + ref.encode() + b'"', data)
            zout.writestr(item, data)
    return str(out)


def test_a_stale_declared_dimension_does_not_cut_the_sheet(store, write_xlsx, tmp_path):
    """openpyxl's read-only mode trusts the sheet's own <dimension> as the
    edge of the data: a workbook declaring A1:B10 but holding 3 columns
    and 101 rows imported 2 columns and 9 rows, ragged 0 — no sign that
    anything was left behind."""
    good = write_xlsx({"Data": [["Time", "User", "Action"]]
                               + [[f"t{i}", f"u{i}", f"a{i}"] for i in range(1, 101)]})
    stale = _with_declared_dimension(good, "A1:B10", tmp_path / "stale.xlsx")

    sheets = store.preview_xlsx_sheets(stale)["tables"]
    assert [c["name"] for c in sheets[0]["columns"]] == ["Time", "User", "Action"]

    rec = store.ingest_xlsx_sheet(stale, "Data", build_fts=False)
    assert rec["row_count"] == 100
    assert [c["name"] for c in rec["columns"]] == ["Time", "User", "Action"]
    v = store.build_view(rec["id"], {"source_id": rec["id"]})
    assert store.fetch_rows(v["view_id"], 99, 1)["rows"][0]["cells"] == ["t100", "u100", "a100"]


def _without_declared_size(tmp_path, rows):
    """A workbook whose one sheet declares no <dimension> — openpyxl's
    write-only mode leaves it out, as many tools' writers do."""
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Log")
    for r in rows:
        ws.append(r)
    path = tmp_path / "undeclared.xlsx"
    wb.save(path)
    return str(path)


def _rows_iter_rows_counts(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        ws = wb.worksheets[0]
        ws.reset_dimensions()
        return sum(1 for _ in ws.iter_rows(values_only=True))
    finally:
        wb.close()


def test_a_sheet_with_no_declared_size_is_counted_without_reading_its_cells(tmp_path, monkeypatch):
    """With no <dimension> to estimate from, the picker counted a sheet's
    rows by reading every cell through openpyxl: 20 s for a 33 MB
    workbook, holding the GIL so a CSV import beside it ran 15x slower,
    and the new import pane asked again each time the file was chosen.
    The count comes off the sheet's row tags now — the same number, across
    several of the 1 MB reads it scans in, with openpyxl's row parser
    reading only the header."""
    from openpyxl.worksheet._reader import WorkSheetParser

    n = 30_000
    path = _without_declared_size(tmp_path, [["When", "Host", "Note"]]
                                  + [[f"2024-05-01 08:{i % 60:02d}:00", f"srv{i % 7}", "x" * 60] for i in range(n)])
    parsed = []
    real = WorkSheetParser.parse_row
    monkeypatch.setattr(WorkSheetParser, "parse_row", lambda self, el: parsed.append(1) or real(self, el))
    sheets = xlsxread.workbook_sheets(path)
    assert sheets == [{"name": "Log", "row_count": n, "columns": ["When", "Host", "Note"]}]
    assert len(parsed) < 5
    monkeypatch.undo()
    assert _rows_iter_rows_counts(path) == n + 1


def test_the_row_count_reads_rows_the_way_openpyxl_does(tmp_path):
    """Prefixed tags (the OpenXML SDK writes <x:row>), rows with no r=""
    (the one before plus one), single quotes, a gap (counted, as iter_rows
    pads it), a cell whose TEXT spells a row tag, and <rowBreaks>: the
    count is iter_rows's, row for row."""
    import zipfile

    path = _without_declared_size(tmp_path, [["placeholder"]])
    cell = '<x:c t="inlineStr"><x:is><x:t>{}</x:t></x:is></x:c>'
    xml = ('<x:worksheet xmlns:x="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><x:sheetData>'
           f'<x:row r="1">{cell.format("Head")}</x:row>'
           f'<x:row>{cell.format("&lt;row r=&quot;900&quot;&gt;")}</x:row>'
           f"<x:row r='5' spans=\"1:1\">{cell.format('five')}</x:row>"
           f'<x:row>{cell.format("six")}</x:row>'
           '</x:sheetData><x:rowBreaks count="1"><x:brk id="40"/></x:rowBreaks></x:worksheet>')
    out = tmp_path / "shaped.xlsx"
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            zout.writestr(item, xml.encode() if item.filename == "xl/worksheets/sheet1.xml" else data)

    assert _rows_iter_rows_counts(str(out)) == 6
    sheet, = xlsxread.workbook_sheets(str(out))
    assert sheet["row_count"] == 5 and sheet["columns"] == ["Head"]


@pytest.mark.parametrize("declared", [True, False], ids=["declared size", "no declared size"])
def test_no_cell_past_the_header_is_cut(store, write_xlsx, tmp_path, monkeypatch, declared):
    """Data wider than the header lands in col_N columns, short rows pad,
    and nothing is counted ragged because nothing was cut. A sheet with no
    <dimension> used to be sized to its header alone, so the cells past it
    were cut and only the ragged count said so."""
    rows = [["When", "Host"],
            ["t1", "h1"],
            ["t2", "h2", "extra", "more"],      # wider than the header
            ["t3"],                             # shorter
            ["t4", "h4", "x", "y", "widest"]]   # wider again, after rows have committed
    path = write_xlsx({"Log": rows}) if declared else _without_declared_size(tmp_path, rows)
    from winnow import store as store_mod
    monkeypatch.setattr(store_mod, "BATCH", 2)   # so the table widens with rows already committed
    rec = store.ingest_xlsx_sheet(path, "Log", build_fts=False)
    assert [c["name"] for c in rec["columns"]] == ["When", "Host", "col_3", "col_4", "col_5"]
    assert rec["ragged_rows"] == 0
    v = store.build_view(rec["id"], {"source_id": rec["id"]})
    assert [r["cells"] for r in store.fetch_rows(v["view_id"], 0, 10)["rows"]] == [
        ["t1", "h1", "", "", ""],
        ["t2", "h2", "extra", "more", ""],
        ["t3", "", "", "", ""],
        ["t4", "h4", "x", "y", "widest"],
    ]
    filt = store.build_view(rec["id"], {"source_id": rec["id"],
                                        "filters": [{"column": "col_5", "op": "equals", "value": "widest"}]})
    assert filt["row_count"] == 1
