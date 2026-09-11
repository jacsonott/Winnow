"""Export all tables to one workbook: every real table, every row, on its
own worksheet — continued on "Name (2)"… past Excel's per-sheet cap, which
the plan route reports first so the UI can ask."""

from __future__ import annotations

import os

from openpyxl import load_workbook

from winnow import log as wlog
from winnow import store as store_mod


def _rows(path, sheet):
    wb = load_workbook(path, read_only=True)
    try:
        return [list(r) for r in wb[sheet].iter_rows(values_only=True)]
    finally:
        wb.close()


def _sheetnames(path):
    wb = load_workbook(path, read_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def test_one_sheet_per_real_table_all_rows_in_rid_order(store, write_csv, tmp_path):
    a = store.ingest_csv(write_csv([["Host", "N"], ["h1", "1"], ["h2", "2"], ["h3", "3"]], "a.csv"))["id"]
    b = store.ingest_csv(write_csv([["Host", "N"], ["x", "9"]], "b.csv"))["id"]
    store.create_merge("both", [a, b])
    tag = store.upsert_tag(None, "hit", "#ff0000", None)["id"]
    store.set_tags(a, [2], tag, True)
    store.set_note(a, 3, "=cmd|' /C calc'!A0")

    out = str(tmp_path / "all.xlsx")
    stats = store.export_all_xlsx(out)
    assert stats == {"sheets": 2, "rows": 4}
    assert _sheetnames(out) == ["a.csv", "b.csv"]          # the merge has no sheet of its own
    rows = _rows(out, "a.csv")
    assert rows[0] == ["Line", "Tags", "Note", "Host", "N"]
    assert [r[0] for r in rows[1:]] == [1, 2, 3]
    assert rows[2][1] == "hit" and rows[1][1] in ("", None)   # an empty cell reads back as None
    assert rows[3][2].startswith("'=cmd"), "formula guard on the note"
    assert rows[1][3:] == ["h1", "1"]


def test_the_layout_decides_columns_and_order(store, write_csv, tmp_path):
    sid = store.ingest_csv(write_csv([["A", "B", "C"], ["1", "2", "3"]], "l.csv"))["id"]
    store.save_layout(sid, {"columns": {"B": {"hidden": True}}, "order": ["C", "A", "B"]})
    out = str(tmp_path / "l.xlsx")
    store.export_all_xlsx(out)
    assert _rows(out, "l.csv")[0] == ["Line", "Tags", "Note", "C", "A"]


def test_a_table_past_the_cap_continues_on_numbered_sheets(store, write_csv, tmp_path, monkeypatch):
    monkeypatch.setattr(store_mod, "XLSX_MAX_ROWS", 6)     # 5 data rows per sheet
    sid = store.ingest_csv(write_csv([["N"]] + [[str(i)] for i in range(12)], "big.csv"))["id"]
    plan = store.export_all_xlsx_plan()
    assert plan["rows_per_sheet"] == 5 and plan["over_cap"] == ["big.csv"]
    (t,) = plan["tables"]
    assert t == {"id": sid, "name": "big.csv", "rows": 12, "sheets": 3, "over_cap": True}

    out = str(tmp_path / "big.xlsx")
    stats = store.export_all_xlsx(out)
    assert stats == {"sheets": 3, "rows": 12}
    assert _sheetnames(out) == ["big.csv", "big.csv (2)", "big.csv (3)"]
    for name, expect in [("big.csv", [1, 2, 3, 4, 5]), ("big.csv (2)", [6, 7, 8, 9, 10]), ("big.csv (3)", [11, 12])]:
        rows = _rows(out, name)
        assert rows[0][0] == "Line", "every continuation repeats the header"
        assert [r[0] for r in rows[1:]] == expect


def test_plan_reports_no_overflow_for_small_tables(store, write_csv):
    store.ingest_csv(write_csv([["A"], ["1"]], "s.csv"))
    plan = store.export_all_xlsx_plan()
    assert plan["over_cap"] == [] and plan["tables"][0]["sheets"] == 1
    assert plan["rows_per_sheet"] == 1_048_575


def test_an_empty_case_still_makes_a_workbook(store, tmp_path):
    out = str(tmp_path / "empty.xlsx")
    assert store.export_all_xlsx(out) == {"sheets": 1, "rows": 0}
    assert _sheetnames(out) == ["No tables"]
    assert store.export_all_xlsx_plan() == {"tables": [], "rows_per_sheet": 1_048_575, "over_cap": []}


def test_the_route_streams_the_file_logs_and_cleans_up(client, store, write_csv):
    wlog.reset()
    store.ingest_csv(write_csv([["A"], ["1"], ["2"]], "r.csv"))
    plan = client.get("/api/export/all_xlsx/plan").json()
    assert plan["tables"][0]["name"] == "r.csv"
    r = client.get("/api/export/all_xlsx")
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert 'filename="all-tables.xlsx"' in r.headers["content-disposition"]
    assert r.content[:2] == b"PK"
    msgs = [e["message"] for e in wlog.entries()]
    assert any(m.startswith("Export started: all tables (.xlsx)") for m in msgs), msgs
    fin = [m for m in msgs if m.startswith("Export finished: all tables (.xlsx)")]
    assert fin and "2 rows on 1 sheet," in fin[0], msgs
    # TestClient runs the BackgroundTask before returning: the temp file is gone.
    import glob
    import tempfile
    assert not glob.glob(os.path.join(tempfile.gettempdir(), "winnow-export-*.xlsx"))


def test_same_named_tables_keep_their_continuations_apart(store, write_csv, tmp_path, monkeypatch):
    """Two hosts' Security.evtx.csv: the second table's first sheet is
    deduped to '…_1', so its overflow must be '…_1 (2)' — not '… (2)',
    which would read as the first table continuing."""
    monkeypatch.setattr(store_mod, "XLSX_MAX_ROWS", 4)     # 3 data rows per sheet
    store.ingest_csv(write_csv([["N"], ["1"], ["2"]], "dup.csv"), name="Security.evtx.csv")
    store.ingest_csv(write_csv([["N"], ["a"], ["b"], ["c"], ["d"], ["e"]], "dup2.csv"), name="Security.evtx.csv")
    out = str(tmp_path / "dup.xlsx")
    store.export_all_xlsx(out)
    assert _sheetnames(out) == ["Security.evtx.csv", "Security.evtx.csv_1", "Security.evtx.csv_1 (2)"]
    assert [r[0] for r in _rows(out, "Security.evtx.csv_1 (2)")[1:]] == [4, 5]
